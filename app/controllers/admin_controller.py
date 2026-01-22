
from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file
from flask_login import login_required, current_user
from app.models.user import User
from app.models.course import Course
from app.models.enrollment import Enrollment
from app.models.course_schedule import CourseSchedule
from app.models.attendance import Attendance
from app import VIDEO_URL, db
from sqlalchemy.orm import joinedload
from datetime import datetime, timedelta
from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
import os

admin_bp = Blueprint('admin', __name__)

def parse_date(date_str, end_of_day=False):
	try:
		dt = datetime.strptime(date_str, '%Y-%m-%d')
		if end_of_day:
			dt = dt + timedelta(hours=23, minutes=59, seconds=59, microseconds=999999)
		return dt
	except (ValueError, TypeError):
		return None

@admin_bp.route('/')
@login_required
def home():
	course_id = request.args.get('course', '').strip()
	student_id = request.args.get('student', '').strip()
	begin_date = parse_date(request.args.get('begin_date', ''))
	end_date = parse_date(request.args.get('end_date', ''), end_of_day=True)

	# Filtros comunes
	filter_conditions = []

	if course_id:
		filter_conditions.append(Attendance.course_id == course_id)
	if student_id:
		filter_conditions.append(Attendance.user_id == student_id)
	if begin_date:
		filter_conditions.append(Attendance.register_date >= begin_date)
	if end_date:
		filter_conditions.append(Attendance.register_date <= end_date)

	if current_user.role == 'admin':
		courses = Course.query.all()
		students = User.query.filter_by(role='student').all()
		attendances = Attendance.query.filter(*filter_conditions).order_by(Attendance.register_date.desc()).all()
		return render_template('dashboard/admin.html', video_url=VIDEO_URL, attendances=attendances, courses=courses, students=students, begin_date=begin_date, end_date=end_date, course_id=course_id, student_id=student_id)
	elif current_user.role == 'student':
		courses = Course.query.join(Enrollment).filter(Enrollment.student_id == current_user.id).all()
		course_ids = [c.id for c in courses]
		filter_conditions.append(Attendance.user_id == current_user.id)
		filter_conditions.append(Attendance.course_id.in_(course_ids))
		attendances = Attendance.query.filter(*filter_conditions).order_by(Attendance.register_date.desc()).all()
		print(attendances)
		return render_template('dashboard/student.html', courses=courses, attendances=attendances, begin_date=begin_date, end_date=end_date, course_id=course_id)
	elif current_user.role == 'teacher':
		courses = Course.query.filter_by(teacher_id=current_user.id).all()
		course_ids = [c.id for c in courses]
		students = User.query.join(Enrollment).filter(Enrollment.course_id.in_(course_ids), User.role == 'student').all()
		filter_conditions.append(Attendance.user_id.in_([s.id for s in students]))
		filter_conditions.append(Attendance.course_id.in_(course_ids))
		attendances = Attendance.query.filter(*filter_conditions).order_by(Attendance.register_date.desc()).all()
		print(courses[0].schedules)
		return render_template('dashboard/teacher.html', courses=courses, attendances=attendances, begin_date=begin_date, end_date=end_date, course_id=course_id, students=students, student_id=student_id)

	return render_template('home.html', user=current_user)

@admin_bp.route('/download-attendance')
@login_required
def download_attendance():
	course_id = request.args.get('course', '').strip()
	student_id = request.args.get('student', '').strip()
	begin_date = parse_date(request.args.get('begin_date', ''))
	end_date = parse_date(request.args.get('end_date', ''), end_of_day=True)

	# Obtener cursos según el rol del usuario
	if current_user.role == 'admin':
		if course_id:
			courses_to_export = Course.query.options(joinedload(Course.teacher)).filter(Course.id == course_id).all()
		else:
			courses_to_export = Course.query.options(joinedload(Course.teacher)).all()
	elif current_user.role == 'student':
		if course_id:
			courses_to_export = Course.query.options(joinedload(Course.teacher)).join(Enrollment).filter(
				Enrollment.student_id == current_user.id,
				Course.id == course_id
			).all()
		else:
			courses_to_export = Course.query.options(joinedload(Course.teacher)).join(Enrollment).filter(
				Enrollment.student_id == current_user.id
			).all()
	elif current_user.role == 'teacher':
		if course_id:
			courses_to_export = Course.query.options(joinedload(Course.teacher)).filter(
				Course.teacher_id == current_user.id,
				Course.id == course_id
			).all()
		else:
			courses_to_export = Course.query.options(joinedload(Course.teacher)).filter(
				Course.teacher_id == current_user.id
			).all()
	else:
		courses_to_export = []

	# Crear el Excel
	wb = Workbook()
	# Eliminar la hoja por defecto
	wb.remove(wb.active)

	# Estilos
	bold_font = Font(bold=True)
	bold_font_large = Font(bold=True, size=14)
	bold_font_medium = Font(bold=True, size=11)
	center_align = Alignment(horizontal='center', vertical='center')
	left_align = Alignment(horizontal='left', vertical='center')
	thin_border = Border(
		left=Side(style='thin'),
		right=Side(style='thin'),
		top=Side(style='thin'),
		bottom=Side(style='thin')
	)
	header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

	# Crear una hoja por cada curso
	for course in courses_to_export:
		# Nombre de la hoja (máximo 31 caracteres)
		sheet_name = course.name[:31] if course.name else "Sin nombre"
		ws = wb.create_sheet(title=sheet_name)

		# Obtener estudiantes matriculados en este curso
		if current_user.role == 'student':
			enrolled_students = [current_user]
		else:
			enrolled_students = User.query.join(Enrollment).filter(
				Enrollment.course_id == course.id,
				User.role == 'student'
			).order_by(User.name).all()

		# Obtener asistencias de este curso
		att_filter = [Attendance.course_id == course.id]
		if student_id:
			att_filter.append(Attendance.user_id == student_id)
		if begin_date:
			att_filter.append(Attendance.register_date >= begin_date)
		if end_date:
			att_filter.append(Attendance.register_date <= end_date)

		attendances = Attendance.query.filter(*att_filter).all()

		# Obtener las fechas únicas donde hay asistencias (solo la fecha, sin hora)
		attendance_dates = set()
		for att in attendances:
			attendance_dates.add(att.register_date.date())

		# Ordenar las fechas
		sorted_dates = sorted(list(attendance_dates))
		num_dates = len(sorted_dates)

		# Calcular el número de columnas para el merge del encabezado
		# Añadimos 2 columnas adicionales para Totales: Asistencias y Faltas
		total_cols = max(3 + num_dates + 2, 20)  # Mínimo 20 columnas para el encabezado
		last_col = get_column_letter(total_cols)

	# Insertar logo institucional si existe (archivo: logo-inst.jpg en app/static)
		logo_path = os.path.join(os.path.dirname(__file__), '..', 'static', 'logo-inst.jpg')
		if os.path.exists(logo_path):
			try:
				logo_img = Image(logo_path)
				# Ajustar tamaño en píxeles según lo deseado
				logo_img.width = 60
				logo_img.height = 70
				# Anclar la imagen en A1 (se sobrepone a las celdas sin cambiar su tamaño)
				try:
					logo_img.anchor = 'A1'
				except Exception:
					pass
				ws.add_image(logo_img, 'A1')
			except Exception:
				pass
				
		# Encabezado institucional
		ws.merge_cells(f'A1:{last_col}1')
		ws['A1'] = "UNIVERSIDAD NACIONAL MICAELA BASTIDAS DE APURÍMAC"
		ws['A1'].font = bold_font_large
		ws['A1'].alignment = center_align

		ws.merge_cells(f'A2:{last_col}2')
		ws['A2'] = "VICERRECTORADO ACADÉMICO"
		ws['A2'].font = bold_font_medium
		ws['A2'].alignment = center_align

		ws.merge_cells(f'A3:{last_col}3')
		ws['A3'] = "DIRECCIÓN DE SERVICIOS ACADÉMICOS – OFICINA DE REGISTRO Y ARCHIVOS"
		ws['A3'].font = bold_font_medium
		ws['A3'].alignment = center_align

		# Información del curso
		row_info = 5
		ws[f'A{row_info}'] = "Escuela Académico Profesional:"
		ws[f'A{row_info}'].font = bold_font
		ws.merge_cells(f'A{row_info}:B{row_info}')
		ws[f'C{row_info}'] = "INGENIERÍA INFORMÁTICA Y SISTEMAS"

		ws[f'L{row_info}'] = "Código de Asignatura:"
		ws[f'L{row_info}'].font = bold_font
		ws[f'N{row_info}'] = course.id[:8].upper() if course.id else "N/A"

		row_info += 1
		ws[f'A{row_info}'] = "Semestre:"
		ws[f'A{row_info}'].font = bold_font
		ws[f'C{row_info}'] = f"SEMESTRE ACADÉMICO {datetime.now().year}-2"

		ws[f'L{row_info}'] = "Creditaje:"
		ws[f'L{row_info}'].font = bold_font
		ws[f'N{row_info}'] = "4"

		row_info += 1
		ws[f'A{row_info}'] = "Asignatura:"
		ws[f'A{row_info}'].font = bold_font
		ws[f'C{row_info}'] = course.name

		ws[f'L{row_info}'] = "Categoría:"
		ws[f'L{row_info}'].font = bold_font
		ws[f'N{row_info}'] = "AFPE"

		row_info += 1
		ws[f'A{row_info}'] = "Docente:"
		ws[f'A{row_info}'].font = bold_font
		ws[f'C{row_info}'] = course.teacher.name.upper() if course.teacher else "N/A"

		ws[f'L{row_info}'] = "Grupo:"
		ws[f'L{row_info}'].font = bold_font
		ws[f'N{row_info}'] = "B"

		# Título de la tabla
		row_title = row_info + 2
		ws.merge_cells(f'A{row_title}:{last_col}{row_title}')
		ws[f'A{row_title}'] = "Lista de Asistencia de Estudiantes"
		ws[f'A{row_title}'].font = bold_font_large
		ws[f'A{row_title}'].alignment = center_align

		# Cabecera de la tabla
		header_row = row_title + 2

		# Cabecera de las columnas fijas
		ws[f'A{header_row}'] = "N°"
		ws[f'B{header_row}'] = "Código Estudiante"
		ws[f'C{header_row}'] = "Apellidos y Nombres"

		for col in ['A', 'B', 'C']:
			ws[f'{col}{header_row}'].font = bold_font
			ws[f'{col}{header_row}'].alignment = center_align
			ws[f'{col}{header_row}'].border = thin_border
			ws[f'{col}{header_row}'].fill = header_fill

		# Agregar columnas de fechas (solo las fechas que tienen registros)
		for i, date in enumerate(sorted_dates):
			col = get_column_letter(4 + i)
			ws[f'{col}{header_row}'] = date.strftime('%d/%m/%Y')
			ws[f'{col}{header_row}'].font = bold_font
			ws[f'{col}{header_row}'].alignment = Alignment(horizontal='center', vertical='center', text_rotation=90)
			ws[f'{col}{header_row}'].border = thin_border
			ws[f'{col}{header_row}'].fill = header_fill

		# Columnas de totales (Asistencias y Faltas)
		col_present = get_column_letter(4 + num_dates)
		col_absent = get_column_letter(4 + num_dates + 1)
		ws[f'{col_present}{header_row}'] = 'Asistencias'
		ws[f'{col_absent}{header_row}'] = 'Faltas'
		for col in [col_present, col_absent]:
			ws[f'{col}{header_row}'].font = bold_font
			ws[f'{col}{header_row}'].alignment = Alignment(horizontal='center', vertical='center')
			ws[f'{col}{header_row}'].border = thin_border
			ws[f'{col}{header_row}'].fill = header_fill

		# Calcular asistencias por fecha para cada estudiante de este curso
		attendance_by_student = {}
		for att in attendances:
			if att.user_id not in attendance_by_student:
				attendance_by_student[att.user_id] = {}

			att_date = att.register_date.date()
			if att_date not in attendance_by_student[att.user_id]:
				attendance_by_student[att.user_id][att_date] = []
			attendance_by_student[att.user_id][att_date].append(att.type)

		# Datos de estudiantes
		data_row = header_row + 1
		for idx, student in enumerate(enrolled_students, 1):
			ws[f'A{data_row}'] = idx
			ws[f'B{data_row}'] = student.id[:6].upper() if student.id else ""
			ws[f'C{data_row}'] = student.name.upper() if student.name else ""

			ws[f'A{data_row}'].alignment = center_align
			ws[f'B{data_row}'].alignment = center_align
			ws[f'C{data_row}'].alignment = left_align
			ws[f'A{data_row}'].border = thin_border
			ws[f'B{data_row}'].border = thin_border
			ws[f'C{data_row}'].border = thin_border

			# Marcar asistencias por fecha
			present_count = 0
			for i, date in enumerate(sorted_dates):
				col = get_column_letter(4 + i)
				cell = ws[f'{col}{data_row}']
				cell.border = thin_border
				cell.alignment = center_align

				if student.id in attendance_by_student and date in attendance_by_student[student.id]:
					cell.value = "A"
					cell.font = Font(color="008000", bold=True)
					present_count += 1
				else:
					cell.value = "F"
					cell.font = Font(color="FF0000", bold=True)
			# Escribir totales de Asistencias y Faltas
			absent_count = num_dates - present_count
			ws[f'{col_present}{data_row}'] = present_count
			ws[f'{col_present}{data_row}'].alignment = center_align
			ws[f'{col_present}{data_row}'].border = thin_border

			ws[f'{col_absent}{data_row}'] = absent_count
			ws[f'{col_absent}{data_row}'].alignment = center_align
			ws[f'{col_absent}{data_row}'].border = thin_border
			
			data_row += 1

		# Ajustar anchos de columna
		ws.column_dimensions['A'].width = 5
		ws.column_dimensions['B'].width = 25
		ws.column_dimensions['C'].width = 50
		for i in range(num_dates):
			ws.column_dimensions[get_column_letter(4 + i)].width = 5

	# Si no hay cursos, crear una hoja vacía
	if not courses_to_export:
		ws = wb.create_sheet(title="Sin datos")
		ws['A1'] = "No hay cursos disponibles para exportar"

	# Guardar en BytesIO
	output = BytesIO()
	wb.save(output)
	output.seek(0)

	# Nombre del archivo
	if course_id and courses_to_export:
		filename = f'asistencia_{courses_to_export[0].name}{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
	else:
		filename = f'asistencia_todos_cursos_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'

	return send_file(
		output,
		as_attachment=True,
		download_name=filename,
		mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
	)

@admin_bp.route('/change-video-url', methods=['POST'])
@login_required
def change_video_url():
	new_url = request.form.get('video_url')
	if current_user.role != 'admin':
		flash('Acción no autorizada', 'error')
		return redirect(url_for('admin.home'))
	global VIDEO_URL
	VIDEO_URL = new_url
	flash('URL de video actualizada exitosamente', 'success')
	return redirect(url_for('admin.home'))